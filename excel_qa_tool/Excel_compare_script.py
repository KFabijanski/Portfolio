from openpyxl import load_workbook
from collections import Counter
import sys
from datetime import datetime, date
import re


EXCLUDED_ROW_PREFIXES = (
    "Data extract produced by",
    "data extract produced by"
)


def normalize_date_token(token):
    token = token.strip()

    # dopasowanie tylko formatów dat, bez ruszania zwykłego tekstu
    formats = [
        "%d.%m.%Y",
        "%m/%d/%Y",
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%d-%m-%Y",
    ]

    for fmt in formats:
        try:
            dt = datetime.strptime(token, fmt)
            return dt.strftime("%Y-%m-%d")
        except ValueError:
            pass

    return token


def normalize_text(value):
    value = value.replace("\xa0", " ")
    value = re.sub(r"\s+", " ", value).strip()
    value = value.casefold()

    value = re.sub(r"\b\d{2}\.\d{2}\.\d{4}\b",
                   lambda m: normalize_date_token(m.group(0)), value)
    value = re.sub(r"\b\d{2}/\d{2}/\d{4}\b",
                   lambda m: normalize_date_token(m.group(0)), value)
    value = re.sub(r"\b\d{4}-\d{2}-\d{2}\b",
                   lambda m: normalize_date_token(m.group(0)), value)
    value = re.sub(r"\b\d{2}-\d{2}-\d{4}\b",
                   lambda m: normalize_date_token(m.group(0)), value)

    return value


def normalize_value(value):
    if value is None:
        return ""

    # prawdziwa data / datetime z openpyxl
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value

    if isinstance(value, str):
        val = normalize_text(value)

        try:
            num = float(val)
            if num.is_integer():
                return int(num)
            return num
        except ValueError:
            return val

    return value


def is_ignored_row(row_values):
    """
    Pomija wiersze techniczne z eksportu.
    """
    non_empty = [v for v in row_values if v != ""]

    if not non_empty:
        return True  # całkiem pusty wiersz

    first = str(non_empty[0]).strip()

    for prefix in EXCLUDED_ROW_PREFIXES:
        if first.startswith(prefix):
            return True

    return False


def normalize_row(row):
    return tuple(normalize_value(cell.value) for cell in row)


def row_to_text(row):
    return " | ".join("empty cell" if v == "" else str(v) for v in row)


def compare_excels(file1, file2):
    wb1 = load_workbook(file1, data_only=True)
    wb2 = load_workbook(file2, data_only=True)

    diffs = []

    sheets1 = set(wb1.sheetnames)
    sheets2 = set(wb2.sheetnames)

    for sheet in sorted(sheets1 - sheets2):
        diffs.append(f'Sheet "{sheet}" exists only in file 1.')
    for sheet in sorted(sheets2 - sheets1):
        diffs.append(f'Sheet "{sheet}" exists only in file 2.')

    for sheet_name in sorted(sheets1 & sheets2):
        ws1 = wb1[sheet_name]
        ws2 = wb2[sheet_name]

        rows1 = []
        rows2 = []

        for row in ws1.iter_rows(values_only=False):
            normalized = normalize_row(row)
            if not is_ignored_row(normalized):
                rows1.append(normalized)

        for row in ws2.iter_rows(values_only=False):
            normalized = normalize_row(row)
            if not is_ignored_row(normalized):
                rows2.append(normalized)

        c1 = Counter(rows1)
        c2 = Counter(rows2)

        only_rows_1 = c1 - c2
        only_rows_2 = c2 - c1

        for row_values, count in only_rows_1.items():
            for _ in range(count):
                diffs.append(f'{sheet_name} - row only in file 1: {row_to_text(row_values)}')

        for row_values, count in only_rows_2.items():
            for _ in range(count):
                diffs.append(f'{sheet_name} - row only in file 2: {row_to_text(row_values)}')

    return diffs


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python compare_excels.py file1.xlsx file2.xlsx")
        sys.exit(1)

    file1 = sys.argv[1]
    file2 = sys.argv[2]

    differences = compare_excels(file1, file2)

    if not differences:
        print("No differences found.")
    else:
        print("Differences found:")
        for diff in differences:
            print(diff)